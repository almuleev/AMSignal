from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "AMSignal_User_Manual_RU.md"
OUTPUT = ROOT / "AMSignal_User_Manual_RU.docx"


def set_font(run, name="Aptos", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(7)


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("AMSignal  Руководство пользователя  Версия 0.15.4")
    set_font(run, size=8)


def add_paragraph_with_terms(document, text):
    paragraph = document.add_paragraph()
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        run = paragraph.add_run(part.strip("`*") if (part.startswith("`") or part.startswith("**")) else part)
        set_font(run, bold=part.startswith("**"))
    return paragraph


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    configure_styles(document)
    add_footer(section)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("# ").strip()
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), name="Aptos Display", size=25)
    document.add_paragraph()
    for line in lines[1:4]:
        if line.strip():
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line.strip()), size=10)
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Версия документа 0.1"), size=10)
    document.add_page_break()

    toc = document.add_paragraph(style="Heading 1")
    toc.add_run("Содержание")
    for entry in (
        "1. О руководстве",
        "2. Быстрый старт",
        "3. Данные и проекты",
        "4. Окно программы и навигация",
        "5. Работа с временными сигналами",
        "6. Измерительные инструменты",
        "7. Преобразование и фильтрация каналов",
        "8. FFT спектральный анализ",
        "9. FRF АЧХ",
        "10. Экспорт и сохранение",
        "11. Настройки и горячие клавиши",
        "12. Практические сценарии",
        "13. Решение проблем и ограничения",
    ):
        p = document.add_paragraph(style="List Bullet")
        set_font(p.add_run(entry))
    document.add_page_break()

    for raw in lines[4:]:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            p = document.add_paragraph(line[3:].strip(), style="Heading 1")
            continue
        if line.startswith("### "):
            p = document.add_paragraph(line[4:].strip(), style="Heading 2")
            continue
        if line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            set_font(p.add_run(line[2:]))
            continue
        if re.match(r"^\d+\. ", line):
            p = document.add_paragraph(style="List Number")
            set_font(p.add_run(re.sub(r"^\d+\. ", "", line)))
            continue
        add_paragraph_with_terms(document, line)

    document.core_properties.title = "Руководство пользователя AMSignal"
    document.core_properties.subject = "Работа с программой AMSignal"
    document.core_properties.author = "AMSignal"
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
